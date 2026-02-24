#!/usr/bin/env python3
"""
generate_presentation.py
Generates NL_Taxonomy_Mapper_V3_Presentation.pptx

Usage:
    python generate_presentation.py

Output:
    NL_Taxonomy_Mapper_V3_Presentation.pptx  (project root)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------------------
# Colour palette (matches GUI)
# ---------------------------------------------------------------------------
C_BLUE   = RGBColor(0x25, 0x63, 0xEB)  # #2563EB primary
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_LGRAY  = RGBColor(0xF1, 0xF5, 0xF9)  # slide backgrounds
C_DARK   = RGBColor(0x1E, 0x29, 0x3B)  # dark text / section headers
C_ACCENT = RGBColor(0x10, 0xB9, 0x81)  # #10B981 green
C_AMBER  = RGBColor(0xF5, 0x9E, 0x0B)
C_PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
C_MUTED  = RGBColor(0x64, 0x74, 0x8B)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _new_prs():
    """Create a 16:9 presentation."""
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank(prs):
    """Add a blank slide (layout 6 — no placeholders)."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_background_fill(slide, color):
    """Fill slide background with a solid colour."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_colored_rect(slide, l, t, w, h, color):
    """Add a solid-filled rectangle with no visible border."""
    shape = slide.shapes.add_shape(1, l, t, w, h)   # 1 = RECTANGLE
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()                      # remove border
    return shape


def add_text_box(slide, text, l, t, w, h,
                 font_size=18, bold=False,
                 color=None, align=PP_ALIGN.LEFT, wrap=True):
    """Add a single-paragraph text box. Returns the shape."""
    txBox = slide.shapes.add_textbox(l, t, w, h)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    return txBox


def add_bullets(slide, items, l, t, w, h):
    """
    items: list of (level, text)
      level 0  →  "•  "  Calibri 15pt  C_DARK
      level 1  →  "  – " Calibri 12pt  C_MUTED
    """
    txBox = slide.shapes.add_textbox(l, t, w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for level, text in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if level == 0:
            prefix, size, col = "\u2022  ", 15, C_DARK
        else:
            prefix, size, col = "    \u2013 ", 12, C_MUTED
        run = p.add_run()
        run.text = prefix + text
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.color.rgb = col
    return txBox


def _set_cell_bg(cell, color):
    """Set table cell background colour via XML."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for elem in tcPr.findall(qn('a:solidFill')):
        tcPr.remove(elem)
    solid = etree.SubElement(tcPr, qn('a:solidFill'))
    srgb  = etree.SubElement(solid, qn('a:srgbClr'))
    srgb.set('val', str(color))


def _set_cell_text(cell, text, font_size=11, bold=False,
                   color=None, align=PP_ALIGN.LEFT):
    cell.text = text
    para = cell.text_frame.paragraphs[0]
    para.alignment = align
    runs = para.runs
    run = runs[0] if runs else para.add_run()
    if not runs:
        run.text = text
    run.font.name  = "Calibri"
    run.font.size  = Pt(font_size)
    run.font.bold  = bold
    if color:
        run.font.color.rgb = color


def add_table(slide, headers, rows, l, t, w, h):
    """
    Styled table: header row C_BLUE/white; alternating white/C_LGRAY rows.
    """
    n_rows = 1 + len(rows)
    n_cols = len(headers)
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, l, t, w, h)
    tbl = tbl_shape.table
    for ci, hdr in enumerate(headers):
        cell = tbl.cell(0, ci)
        _set_cell_bg(cell, C_BLUE)
        _set_cell_text(cell, hdr, font_size=11, bold=True, color=C_WHITE)
    for ri, row_data in enumerate(rows):
        bg = C_WHITE if ri % 2 == 0 else C_LGRAY
        for ci, val in enumerate(row_data):
            cell = tbl.cell(ri + 1, ci)
            _set_cell_bg(cell, bg)
            _set_cell_text(cell, str(val), font_size=11, color=C_DARK)
    return tbl_shape


# ---------------------------------------------------------------------------
# Composite slide-template helpers
# ---------------------------------------------------------------------------

def add_section_header(prs, title, subtitle=""):
    """Full-bleed section divider: C_DARK bg, white 36pt title, accent subtitle."""
    slide = _blank(prs)
    add_background_fill(slide, C_DARK)
    add_colored_rect(slide, Inches(0), Inches(2.8), Inches(0.18), Inches(1.9), C_BLUE)
    add_text_box(slide, title,
                 Inches(0.5), Inches(2.8), Inches(12.0), Inches(1.2),
                 font_size=36, bold=True, color=C_WHITE)
    if subtitle:
        add_text_box(slide, subtitle,
                     Inches(0.5), Inches(4.0), Inches(12.0), Inches(0.6),
                     font_size=20, color=C_ACCENT)
    return slide


def add_content_slide(prs, title):
    """
    C_LGRAY background + white card + C_BLUE left bar + slide title.
    Returns the ready slide.
    """
    slide = _blank(prs)
    add_background_fill(slide, C_LGRAY)
    add_colored_rect(slide, Inches(0.3),  Inches(0.3),  Inches(12.73), Inches(6.9),  C_WHITE)
    add_colored_rect(slide, Inches(0.3),  Inches(0.3),  Inches(0.10),  Inches(6.9),  C_BLUE)
    add_text_box(slide, title,
                 Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.7),
                 font_size=24, bold=True, color=C_DARK)
    add_colored_rect(slide, Inches(0.6), Inches(1.05), Inches(12.3), Inches(0.03), C_LGRAY)
    return slide


# ---------------------------------------------------------------------------
# Slide 1 — Title
# ---------------------------------------------------------------------------
def slide_01_title(prs):
    print("Slide  1/19: Title")
    slide = _blank(prs)
    add_background_fill(slide, C_DARK)
    add_text_box(slide, "Taxonomy Mapper V3",
                 Inches(0.7), Inches(1.2), Inches(11.5), Inches(1.4),
                 font_size=48, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, "Map content URLs to taxonomy topics using fuzzy matching",
                 Inches(0.7), Inches(2.7), Inches(11.5), Inches(0.7),
                 font_size=22, color=C_ACCENT, align=PP_ALIGN.CENTER)
    add_text_box(slide, "NL  \u00b7  BE  \u00b7  GB  \u00b7  SE",
                 Inches(0.7), Inches(3.35), Inches(11.5), Inches(0.5),
                 font_size=18, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_colored_rect(slide, Inches(1.5), Inches(3.5), Inches(10.33), Inches(0.06), C_BLUE)
    add_text_box(slide, "v3.34  \u00b7  Feb 2026",
                 Inches(9.5), Inches(6.9), Inches(3.5), Inches(0.4),
                 font_size=13, color=C_MUTED, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------------------
# Slide 2 — Problem Statement
# ---------------------------------------------------------------------------
def slide_02_problem(prs):
    print("Slide  2/19: Problem Statement")
    slide = add_content_slide(prs, "The Challenge: Mapping Content at Scale")
    items = [
        (0, "Manual content tagging is slow, inconsistent, and doesn't scale"),
        (0, "Content gaps in taxonomy coverage are invisible without analysis"),
        (0, "Taxonomy Mapper automates fuzzy keyword-to-topic matching"),
        (1, "Processes hundreds of URLs in seconds — not days"),
        (1, "Supports Dutch, French, English, Swedish content"),
        (0, "Confidence scoring makes match quality visible at a glance"),
        (1, "Best Match \u2192 Highly Relevant \u2192 Relevant \u2192 Tangential \u2192 Unmapped"),
        (0, "Multi-country taxonomy management: NL, BE, GB, SE"),
        (0, "Integrated tools: Campaign Assistant, Gap Analysis, Synonym Editor"),
    ]
    add_bullets(slide, items, Inches(0.7), Inches(1.2), Inches(12.0), Inches(5.8))


# ---------------------------------------------------------------------------
# Slide 3 — How It Works
# ---------------------------------------------------------------------------
def slide_03_how_it_works(prs):
    print("Slide  3/19: How It Works")
    slide = add_content_slide(prs, "How It Works")
    items = [
        (0, "Semantic carriers + taxonomy file as inputs"),
        (0, "Fuzzy matching engine powered by rapidfuzz"),
        (1, "50\u00d7 faster than fuzzywuzzy, same API"),
        (0, "Synonyms bridge language & terminology gaps"),
        (1, "Bidirectional expansion, word-boundary enforced"),
        (0, "Deduplication: one row per URL\u2013Segment"),
        (0, "Threshold (default 80%) controls match sensitivity"),
        (0, "Output auto-saves to Excel with relevance labels"),
    ]
    add_bullets(slide, items, Inches(0.7), Inches(1.2), Inches(5.8), Inches(5.5))

    # Right: horizontal flow diagram
    flow = [
        ("Keywords",         "Keyword 1-12\n+ Summary"),
        ("Synonym\nExpansion", "Bidirectional\nword-boundary"),
        ("Fuzzy\nMatch",     "rapidfuzz\nratio()"),
        ("Score &\nRank",    "Relevance\nlabels"),
        ("Output\nExcel",    "One row per\nURL\u2013Segment"),
    ]
    box_w = Inches(1.55)
    box_h = Inches(1.1)
    gap   = Inches(0.22)
    x     = Inches(6.9)
    y     = Inches(2.6)
    for i, (label, sub) in enumerate(flow):
        add_colored_rect(slide, x, y, box_w, box_h, C_BLUE)
        add_text_box(slide, label, x, y + Inches(0.08),
                     box_w, Inches(0.5), font_size=11, bold=True,
                     color=C_WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, sub, x, y + Inches(0.55),
                     box_w, Inches(0.5), font_size=9,
                     color=C_WHITE, align=PP_ALIGN.CENTER)
        if i < len(flow) - 1:
            add_text_box(slide, "\u2192", x + box_w, y + Inches(0.3),
                         gap, Inches(0.4), font_size=14,
                         color=C_DARK, align=PP_ALIGN.CENTER)
        x += box_w + gap


# ---------------------------------------------------------------------------
# Slide 4 — Supported Countries
# ---------------------------------------------------------------------------
def slide_04_countries(prs):
    print("Slide  4/19: Multi-Country Support")
    slide = add_content_slide(prs, "Multi-Country Support")
    countries = [
        ("NL", "Netherlands",    "Dutch (NL)"),
        ("BE", "Belgium",        "Dutch / French"),
        ("GB", "United Kingdom", "English (EN)"),
        ("SE", "Sweden",         "Swedish (SE)"),
    ]
    positions = [
        (Inches(0.7),  Inches(1.3)),
        (Inches(6.7),  Inches(1.3)),
        (Inches(0.7),  Inches(3.55)),
        (Inches(6.7),  Inches(3.55)),
    ]
    card_w = Inches(5.5)
    card_h = Inches(2.0)
    for (code, name, lang), (cx, cy) in zip(countries, positions):
        add_colored_rect(slide, cx, cy, card_w, card_h, C_WHITE)
        add_colored_rect(slide, cx, cy, Inches(0.12), card_h, C_BLUE)
        badge_w, badge_h = Inches(0.75), Inches(0.42)
        add_colored_rect(slide, cx + Inches(0.22), cy + Inches(0.15),
                         badge_w, badge_h, C_BLUE)
        add_text_box(slide, code,
                     cx + Inches(0.22), cy + Inches(0.15),
                     badge_w, badge_h,
                     font_size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, name,
                     cx + Inches(1.1), cy + Inches(0.15),
                     Inches(4.0), Inches(0.5),
                     font_size=16, bold=True, color=C_DARK)
        add_text_box(slide, lang,
                     cx + Inches(1.1), cy + Inches(0.68),
                     Inches(4.0), Inches(0.4),
                     font_size=13, color=C_MUTED)


# ---------------------------------------------------------------------------
# Slide 5 — End-to-End Workflow
# ---------------------------------------------------------------------------
def slide_05_workflow(prs):
    print("Slide  5/19: End-to-End Workflow")
    slide = add_content_slide(prs, "End-to-End Workflow")
    steps = [
        (C_BLUE,   "Prepare",  "Load semantic\ncarriers + taxonomy"),
        (C_BLUE,   "Match",    "Run Matching \u2192\noutput auto-saved"),
        (C_ACCENT, "Review",   "Campaign Assistant\n\u2192 Analyse"),
        (C_ACCENT, "Improve",  "Apply category /\nsynonym fixes"),
        (C_BLUE,   "Re-run",   "Compare\nbefore/after delta"),
    ]
    box_w  = Inches(2.1)
    box_h  = Inches(2.6)
    conn_w = Inches(0.25)
    x      = Inches(0.65)
    y      = Inches(1.9)
    for i, (color, label, sub) in enumerate(steps):
        add_colored_rect(slide, x, y, box_w, box_h, color)
        add_text_box(slide, str(i + 1),
                     x, y + Inches(0.12), box_w, Inches(0.45),
                     font_size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, label,
                     x, y + Inches(0.6), box_w, Inches(0.55),
                     font_size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, sub,
                     x, y + Inches(1.2), box_w, Inches(1.1),
                     font_size=11, color=C_WHITE, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_colored_rect(slide, x + box_w, y + Inches(1.1),
                             conn_w, Inches(0.06), C_MUTED)
        x += box_w + conn_w


# ---------------------------------------------------------------------------
# Slide 6 — Input Files
# ---------------------------------------------------------------------------
def slide_06_input_files(prs):
    print("Slide  6/19: Input Files")
    slide = add_content_slide(prs, "Input Files")

    # Left column header
    add_colored_rect(slide, Inches(0.65), Inches(1.2), Inches(5.7), Inches(0.45), C_BLUE)
    add_text_box(slide, "Semantic Carriers (URL Keywords File)",
                 Inches(0.75), Inches(1.2), Inches(5.5), Inches(0.45),
                 font_size=12, bold=True, color=C_WHITE)
    add_bullets(slide, [
        (0, "URL column (required)"),
        (0, "Keyword 1 \u2013 Keyword 12 (required)"),
        (0, "Title / Summary / Description (optional)"),
        (0, "Product column (optional \u2014 ignored in matching)"),
        (0, "Salesforce: Data_Categories__c \u2192 multi-product matching"),
    ], Inches(0.7), Inches(1.75), Inches(5.5), Inches(2.5))

    # Right column header
    add_colored_rect(slide, Inches(6.9), Inches(1.2), Inches(5.7), Inches(0.45), C_BLUE)
    add_text_box(slide, "Taxonomy File",
                 Inches(7.0), Inches(1.2), Inches(5.5), Inches(0.45),
                 font_size=12, bold=True, color=C_WHITE)
    add_bullets(slide, [
        (0, "Product \u2192 Domain \u2192 Segment \u2192 Topic columns"),
        (0, "Topics detected dynamically (Topic 1 \u2026 Topic N)"),
        (0, "Blank Product = matches any semantic Product"),
        (0, "restructure_taxonomy.py converts multi-product columns"),
    ], Inches(6.95), Inches(1.75), Inches(5.5), Inches(2.5))

    # Bottom table
    add_text_box(slide, "Key Output Columns",
                 Inches(0.7), Inches(4.3), Inches(5.5), Inches(0.4),
                 font_size=13, bold=True, color=C_DARK)
    add_table(slide,
              ["Column", "Description"],
              [
                  ("URL",                    "Source content URL"),
                  ("Product / Domain / Segment", "Taxonomy hierarchy"),
                  ("Topic_1 \u2026 Topic_N",  "Matched topics (up to 6)"),
                  ("Top_Score",               "Highest match score 0\u2013100"),
                  ("Top_Relevance",           "Relevance label"),
                  ("Unmapped_Reason",         "Why no match (blank if matched)"),
              ],
              Inches(0.7), Inches(4.7), Inches(12.0), Inches(2.0))


# ---------------------------------------------------------------------------
# Slide 7 — Running the Matcher
# ---------------------------------------------------------------------------
def slide_07_running_matcher(prs):
    print("Slide  7/19: Running the Matcher")
    slide = add_content_slide(prs, "Setup Tab: Running the Matcher")
    steps = [
        ("1", "Select country (NL / BE / GB / SE)"),
        ("2", "Browse to semantic carriers file (URL + Keyword 1\u201312)"),
        ("3", "Browse to taxonomy file (auto-populated from last run)"),
        ("4", "Set similarity threshold (default 80%)"),
        ("5", 'Enable \u201cUse Summary\u201d for Summary column matching (optional)'),
        ("6", 'Click \u201cRun Matching\u201d \u2192 taxonomy_match_{CC}_{timestamp}.xlsx'),
    ]
    badge_w, badge_h = Inches(0.4), Inches(0.4)
    x_badge = Inches(0.7)
    x_text  = Inches(1.3)
    y       = Inches(1.25)
    step_h  = Inches(0.72)
    for num, text in steps:
        add_colored_rect(slide, x_badge, y, badge_w, badge_h, C_BLUE)
        add_text_box(slide, num, x_badge, y, badge_w, badge_h,
                     font_size=13, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, text, x_text, y + Inches(0.02),
                     Inches(11.0), Inches(0.4), font_size=14, color=C_DARK)
        y += step_h
    # Amber callout
    add_colored_rect(slide, Inches(0.7), Inches(6.1), Inches(11.5), Inches(0.7), C_AMBER)
    add_text_box(slide,
                 "\u26a0  File swap detection: if files appear reversed, click Swap Files",
                 Inches(0.9), Inches(6.12), Inches(11.2), Inches(0.5),
                 font_size=13, bold=True, color=C_WHITE)


# ---------------------------------------------------------------------------
# Slide 8 — Understanding the Output
# ---------------------------------------------------------------------------
def slide_08_output(prs):
    print("Slide  8/19: Understanding the Output")
    slide = add_content_slide(prs, "Understanding the Output")
    add_text_box(slide,
                 "One row per URL\u2013Segment. Topics spread across Topic_1 \u2026 Topic_N (up to 6 per row).",
                 Inches(0.7), Inches(1.15), Inches(12.0), Inches(0.45),
                 font_size=13, color=C_MUTED)
    add_table(slide,
              ["Column", "Description"],
              [
                  ("URL",                    "Source content URL"),
                  ("Product / Domain / Segment", "Taxonomy hierarchy from matched topic"),
                  ("Topic_1 \u2026 Topic_N",  "All matched topics \u2014 up to 6 per URL\u2013Segment"),
                  ("Top_Score",               "Highest fuzzy match score (0\u2013100)"),
                  ("Top_Relevance",           "Relevance label (Best Match \u2192 Unmapped)"),
                  ("Unmapped_Reason",         "Why no matches (blank on matched rows)"),
                  ("Suggested_Product",       "Pre-filled for product-filtered rows only"),
                  ("Unmatched_Keywords",      "Keywords that scored below threshold"),
              ],
              Inches(0.7), Inches(1.65), Inches(12.0), Inches(5.1))


# ---------------------------------------------------------------------------
# Slide 9 — Match Quality Levels
# ---------------------------------------------------------------------------
def slide_09_quality(prs):
    print("Slide  9/19: Match Quality Levels")
    slide = add_content_slide(prs, "Match Quality Levels")
    rows_data = [
        ("Best Match",        "Score \u2265 90%  AND  topic keyword found in URL",   C_ACCENT),
        ("Highly Relevant",   "Score \u2265 85%  AND  topic keyword found in URL",   C_ACCENT),
        ("Relevant",          "Score < 85%  BUT  topic keyword found in URL",         C_BLUE),
        ("Somewhat Relevant", "Score \u2265 90%,  topic NOT found in URL",            C_BLUE),
        ("Tangential",        "Score \u2265 85%,  topic NOT found in URL",            C_MUTED),
        ("Low Relevance",     "Score 80\u201384%, topic NOT found in URL",            C_AMBER),
        ("Unmapped",          "No matches above threshold (score = 0)",               C_AMBER),
    ]
    n_rows = 1 + len(rows_data)
    tbl_shape = slide.shapes.add_table(n_rows, 2,
                                       Inches(0.7), Inches(1.3),
                                       Inches(12.0), Inches(5.8))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(3.5)
    tbl.columns[1].width = Inches(8.5)
    # Header
    for ci, hdr in enumerate(["Relevance Level", "Criteria"]):
        cell = tbl.cell(0, ci)
        _set_cell_bg(cell, C_DARK)
        _set_cell_text(cell, hdr, font_size=12, bold=True, color=C_WHITE)
    # Data rows
    for ri, (label, criteria, badge_color) in enumerate(rows_data):
        bg = C_WHITE if ri % 2 == 0 else C_LGRAY
        cell0 = tbl.cell(ri + 1, 0)
        cell1 = tbl.cell(ri + 1, 1)
        _set_cell_bg(cell0, badge_color)
        _set_cell_bg(cell1, bg)
        _set_cell_text(cell0, label, font_size=11, bold=True, color=C_WHITE)
        _set_cell_text(cell1, criteria, font_size=11, color=C_DARK)


# ---------------------------------------------------------------------------
# Slide 10 — Post-Run Review Dialog
# ---------------------------------------------------------------------------
def slide_10_post_run(prs):
    print("Slide 10/19: Post-Run Review Dialog")
    slide = add_content_slide(prs, "Post-Run Review: Automatic Synonym Suggestions")
    add_bullets(slide, [
        (0, "Appears automatically after every matching run"),
        (0, "Shows HIGH and MEDIUM priority near-misses only"),
        (1, "HIGH: score \u2265 threshold\u22125, frequency \u2265 3"),
        (1, "MEDIUM: score \u2265 threshold\u221215, frequency \u2265 2"),
        (1, "LOW: score \u2265 50 (hidden by default)"),
        (0, "Per-row Approve / Dismiss actions"),
        (0, "\u201cShow URLs\u201d expands URL list for each suggestion"),
        (0, "Approve writes directly to synonyms.json with backup"),
    ], Inches(0.7), Inches(1.2), Inches(6.3), Inches(5.0))

    # Callout box
    add_colored_rect(slide, Inches(7.3), Inches(1.8), Inches(5.5), Inches(3.0), C_LGRAY)
    add_colored_rect(slide, Inches(7.3), Inches(1.8), Inches(5.5), Inches(0.1),  C_BLUE)
    add_text_box(slide,
                 "This is the primary\nsynonym workflow\n\nUse it after every run\nto capture near-misses",
                 Inches(7.5), Inches(2.0), Inches(5.1), Inches(2.4),
                 font_size=16, color=C_DARK, align=PP_ALIGN.CENTER)

    # Priority table
    add_table(slide,
              ["Priority", "Score Threshold", "Min Frequency"],
              [
                  ("HIGH",   "\u2265 threshold \u2212 5",  "3+"),
                  ("MEDIUM", "\u2265 threshold \u2212 15", "2+"),
                  ("LOW",    "\u2265 50",                  "1+"),
              ],
              Inches(7.3), Inches(5.0), Inches(5.5), Inches(1.6))


# ---------------------------------------------------------------------------
# Slide 11 — Gap Analysis Report
# ---------------------------------------------------------------------------
def slide_11_gap_analysis(prs):
    print("Slide 11/19: Gap Analysis Report")
    slide = add_content_slide(prs, "Gap Analysis Report: Checking Match Rate")
    add_bullets(slide, [
        (0, "Reports tab \u2192 Gap Analysis button"),
        (0, "Generates an 8-sheet Excel workbook"),
        (0, "Per-unique-URL match rate (authoritative metric)"),
        (1, "Console log count is per-row \u2014 inflated by multi-product duplicates"),
        (1, "For BE these can differ by \u223c3 pp"),
        (0, "Tracks rate across multiple runs for trend analysis"),
        (0, "Identifies phantom topics and never-matched topics"),
    ], Inches(0.7), Inches(1.2), Inches(6.0), Inches(4.5))

    # 8 sheets list
    add_colored_rect(slide, Inches(7.2), Inches(1.2), Inches(5.6), Inches(0.45), C_DARK)
    add_text_box(slide, "8 Report Sheets",
                 Inches(7.3), Inches(1.2), Inches(5.4), Inches(0.45),
                 font_size=13, bold=True, color=C_WHITE)
    y = Inches(1.72)
    for name in [
        "1. Summary",
        "2. Never-Matched Topics",
        "3. Phantom Topics",
        "4. Synonym Recommendations",
        "5. URL Coverage",
        "6. Topic Frequency",
        "7. Match Rate Trend",
        "8. Full Data",
    ]:
        add_text_box(slide, name,
                     Inches(7.3), y, Inches(5.4), Inches(0.38),
                     font_size=12, color=C_DARK)
        y += Inches(0.43)

    # Callout
    add_colored_rect(slide, Inches(0.65), Inches(5.85), Inches(12.0), Inches(0.8), C_ACCENT)
    add_text_box(slide,
                 "Always use Gap Analysis (not the console log) for match rate. "
                 "Console count can be \u223c3 pp higher for BE (multi-product inflation).",
                 Inches(0.85), Inches(5.9), Inches(11.7), Inches(0.65),
                 font_size=12, bold=True, color=C_WHITE)


# ---------------------------------------------------------------------------
# Slide 12 — Campaign Assistant
# ---------------------------------------------------------------------------
def slide_12_campaign_assistant(prs):
    print("Slide 12/19: Campaign Assistant")
    slide = add_content_slide(prs, "Campaign Assistant: Automated Improvement Advisor")

    # Workflow strip
    flow  = ["Load Match File", "Auto-fill + Analyse", "Apply Suggestions"]
    sw    = Inches(3.7)
    sh    = Inches(0.5)
    sx    = Inches(0.7)
    sy    = Inches(1.2)
    for i, label in enumerate(flow):
        add_colored_rect(slide, sx, sy, sw, sh, C_BLUE)
        add_text_box(slide, label, sx, sy, sw, sh,
                     font_size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        sx += sw
        if i < len(flow) - 1:
            add_text_box(slide, "\u2192", sx, sy, Inches(0.3), sh,
                         font_size=16, color=C_DARK, align=PP_ALIGN.CENTER)
            sx += Inches(0.3)

    # 2×2 suggestion type cards
    cards = [
        (C_BLUE,   "CATEGORY MAP",
         "Salesforce product tag missing from\ncategory_mapping.json"),
        (C_ACCENT, "SYNONYM",
         "Near-miss keyword \u2014 add synonym\nto bridge the terminology gap"),
        (C_AMBER,  "DATA ISSUE",
         "No extractable keywords \u2014 check\nTitle / Summary columns"),
        (C_PURPLE, "NO MATCH",
         "Keywords exist but nothing scores\nabove threshold"),
    ]
    positions = [
        (Inches(0.7), Inches(2.0)),
        (Inches(6.9), Inches(2.0)),
        (Inches(0.7), Inches(3.8)),
        (Inches(6.9), Inches(3.8)),
    ]
    card_w, card_h = Inches(5.7), Inches(1.6)
    for (color, badge, desc), (cx, cy) in zip(cards, positions):
        add_colored_rect(slide, cx, cy, card_w, card_h, C_WHITE)
        add_colored_rect(slide, cx, cy, Inches(0.15), card_h, color)
        add_colored_rect(slide, cx + Inches(0.22), cy + Inches(0.18),
                         Inches(1.8), Inches(0.38), color)
        add_text_box(slide, badge,
                     cx + Inches(0.22), cy + Inches(0.18),
                     Inches(1.8), Inches(0.38),
                     font_size=10, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, desc,
                     cx + Inches(0.22), cy + Inches(0.65),
                     Inches(5.3), Inches(0.85),
                     font_size=12, color=C_DARK)

    add_text_box(slide,
                 "Each card shows estimated URL gain and allows URL drill-down.",
                 Inches(0.7), Inches(5.55), Inches(12.0), Inches(0.4),
                 font_size=12, color=C_MUTED)


# ---------------------------------------------------------------------------
# Slide 13 — Section Header (Improvement Loop)
# ---------------------------------------------------------------------------
def slide_section_improvement(prs):
    print("Slide 13/19: Section Header \u2014 Improvement Loop")
    add_section_header(prs, "Improvement Loop", "Close the gap systematically")


# ---------------------------------------------------------------------------
# Slide 14 — The Improvement Loop
# ---------------------------------------------------------------------------
def slide_14_improvement_loop(prs):
    print("Slide 14/19: The Improvement Loop")
    slide = add_content_slide(prs, "The Improvement Loop")
    steps = [
        ("1", "Load Files (Setup)",     "Select semantic carriers + taxonomy \u2192 Run"),
        ("2", "Run the Matcher",         "Output: taxonomy_match_{CC}_{timestamp}.xlsx"),
        ("3", "Check Rate & Diagnose",   "Campaign Assistant \u2192 Auto-fill \u2192 Analyse"),
        ("4", "Apply Fixes",             "Category map entries or synonyms from suggestion cards"),
        ("5", "Re-run Matcher",          "Before/after delta shown in Campaign Assistant"),
        ("6", "Generate Gap Analysis",   "Reports \u2192 Gap Analysis for full coverage view"),
    ]
    badge_w, badge_h = Inches(0.45), Inches(0.45)
    y = Inches(1.25)
    for num, title, desc in steps:
        add_colored_rect(slide, Inches(0.7), y, badge_w, badge_h, C_BLUE)
        add_text_box(slide, num, Inches(0.7), y, badge_w, badge_h,
                     font_size=13, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, title,
                     Inches(1.35), y, Inches(11.0), Inches(0.38),
                     font_size=14, bold=True, color=C_DARK)
        add_text_box(slide, desc,
                     Inches(1.35), y + Inches(0.36), Inches(11.0), Inches(0.35),
                     font_size=12, color=C_MUTED)
        y += Inches(0.88)


# ---------------------------------------------------------------------------
# Slide 15 — Synonym Management: 3 Methods
# ---------------------------------------------------------------------------
def slide_15_synonyms(prs):
    print("Slide 15/19: Adding Synonyms: 3 Methods")
    slide = add_content_slide(prs, "Adding Synonyms: 3 Methods")
    methods = [
        (C_ACCENT, "AUTOMATIC",
         "Post-run dialog\nappears after every run",
         "\u2022 Per-row Approve / Dismiss\n\u2022 Best for daily use\n\u2022 Direct write to synonyms.json\n\u2022 Backup created automatically"),
        (C_BLUE,   "MANUAL",
         "Synonym Editor tab",
         "\u2022 Bulk import via clipboard\n\u2022 Import from Taxonomy\n\u2022 Double-click to edit inline\n\u2022 Best for targeted additions"),
        (C_PURPLE, "BROWSER TOOL",
         "synonym_review.html",
         "\u2022 Reads Keyword Recommendations\n\u2022 Multi-select batch export\n\u2022 python apply_synonym_patch.py\n\u2022 Best for bulk discovery"),
    ]
    card_w, card_h = Inches(3.9), Inches(5.0)
    x, y = Inches(0.7), Inches(1.2)
    for color, badge, subtitle, bullets in methods:
        add_colored_rect(slide, x, y, card_w, card_h, C_WHITE)
        add_colored_rect(slide, x, y, card_w, Inches(0.08), color)
        add_colored_rect(slide, x + Inches(0.12), y + Inches(0.18),
                         Inches(1.8), Inches(0.38), color)
        add_text_box(slide, badge,
                     x + Inches(0.12), y + Inches(0.18),
                     Inches(1.8), Inches(0.38),
                     font_size=10, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_text_box(slide, subtitle,
                     x + Inches(0.2), y + Inches(0.72),
                     Inches(3.5), Inches(0.5),
                     font_size=13, bold=True, color=C_DARK)
        add_text_box(slide, bullets,
                     x + Inches(0.2), y + Inches(1.3),
                     Inches(3.55), Inches(3.4),
                     font_size=12, color=C_DARK)
        x += card_w + Inches(0.35)


# ---------------------------------------------------------------------------
# Slide 16 — Unmapped Review Tool
# ---------------------------------------------------------------------------
def slide_16_unmapped_review(prs):
    print("Slide 16/19: Unmapped Review Tool")
    slide = add_content_slide(prs, "Unmapped Review Tool (unmapped_review.html)")
    add_text_box(slide,
                 "Row-by-row triage for URLs that bulk methods can't resolve. "
                 "Open in a browser \u2014 loads match output directly.",
                 Inches(0.7), Inches(1.15), Inches(12.0), Inches(0.5),
                 font_size=13, color=C_MUTED)
    add_table(slide,
              ["Action", "When to Use", "Export File"],
              [
                  ("Remove",  "Genuinely off-topic URL \u2014 should not be in scope",    "unmapped_exclusions_*.json"),
                  ("Product", "Wrong product assigned \u2014 correct the category mapping", "category_fixes_*.json"),
                  ("Synonym", "Near-miss keyword \u2014 add synonym to bridge the gap",    "synonym_patch_*.json"),
                  ("Noise",   "CMS boilerplate keywords polluting extraction",              "noise_phrases_*.txt"),
                  ("Skip",    "Uncertain \u2014 mark for later review",                    "(none)"),
              ],
              Inches(0.7), Inches(1.75), Inches(12.0), Inches(3.5))

    add_colored_rect(slide, Inches(0.7), Inches(5.4), Inches(12.0), Inches(0.75), C_BLUE)
    add_text_box(slide,
                 "Tip: Work tabs in order \u2014 Product (highest yield) \u2192 No Match \u2192 Artikel Noise",
                 Inches(0.9), Inches(5.43), Inches(11.7), Inches(0.6),
                 font_size=13, bold=True, color=C_WHITE)


# ---------------------------------------------------------------------------
# Slide 17 — Reports Overview
# ---------------------------------------------------------------------------
def slide_17_reports(prs):
    print("Slide 17/19: Available Reports")
    slide = add_content_slide(prs, "Available Reports")

    # Left: Quick Reports
    add_colored_rect(slide, Inches(0.65), Inches(1.2), Inches(5.7), Inches(0.45), C_DARK)
    add_text_box(slide, "Quick Reports (Reports tab)",
                 Inches(0.75), Inches(1.2), Inches(5.5), Inches(0.45),
                 font_size=12, bold=True, color=C_WHITE)
    y = Inches(1.75)
    for color, name, desc in [
        (C_BLUE,   "Proposed Synonyms",  "5-sheet Excel \u00b7 HIGH/MEDIUM/LOW ranked suggestions"),
        (C_PURPLE, "Match Quality",       "5-sheet Excel \u00b7 relevance distribution, low-confidence rows"),
        (C_AMBER,  "Unmapped Reasons",    "5-sheet Excel \u00b7 failure diagnostics by reason type"),
    ]:
        add_colored_rect(slide, Inches(0.7), y, Inches(0.12), Inches(0.65), color)
        add_text_box(slide, name,
                     Inches(0.9), y, Inches(5.2), Inches(0.35),
                     font_size=13, bold=True, color=C_DARK)
        add_text_box(slide, desc,
                     Inches(0.9), y + Inches(0.32), Inches(5.2), Inches(0.35),
                     font_size=11, color=C_MUTED)
        y += Inches(0.85)

    # Right: Full Reports
    add_colored_rect(slide, Inches(6.9), Inches(1.2), Inches(5.7), Inches(0.45), C_DARK)
    add_text_box(slide, "Full Reports",
                 Inches(7.0), Inches(1.2), Inches(5.5), Inches(0.45),
                 font_size=12, bold=True, color=C_WHITE)
    y = Inches(1.75)
    for name, desc in [
        ("Gap Analysis",
         "8 sheets \u00b7 phantom/never-matched topics,\ncoverage trends, synonym recommendations"),
        ("Topic Recommendations",
         "8 sheets \u00b7 content gaps vs taxonomy demand,\nnew topic proposals"),
    ]:
        add_colored_rect(slide, Inches(6.95), y, Inches(0.12), Inches(0.9), C_ACCENT)
        add_text_box(slide, name,
                     Inches(7.15), y, Inches(5.2), Inches(0.38),
                     font_size=13, bold=True, color=C_DARK)
        add_text_box(slide, desc,
                     Inches(7.15), y + Inches(0.35), Inches(5.2), Inches(0.5),
                     font_size=11, color=C_MUTED)
        y += Inches(1.1)

    add_text_box(slide,
                 "Note: One-Click Apply available after Proposed Synonyms report. "
                 "Post-run dialog remains the primary synonym workflow.",
                 Inches(0.7), Inches(6.15), Inches(12.0), Inches(0.55),
                 font_size=12, color=C_MUTED)


# ---------------------------------------------------------------------------
# Slide 18 — Advanced Features
# ---------------------------------------------------------------------------
def slide_18_advanced(prs):
    print("Slide 18/19: Advanced Features")
    slide = add_content_slide(prs, "Advanced Features")
    features = [
        (C_BLUE,   "URL Pattern Filter",
         "Remove nav/special/edit-mode pages before matching",
         "GUI: Setup \u2192 Filter URL Patterns\nCLI: --url-filter pattern\n\nFilters full URLs via regex; runs before keyword extraction so filtered pages never appear in output."),
        (C_ACCENT, "Content Keywords\nExtractor",
         "Crawl live URLs for Title / Description / Summary",
         "Requires: pip install trafilatura\nCLI: python content_keyword_extractor.py\n\nDeduplicates by URL (unless Product column present). Appends crawled content to semantic carrier file."),
        (C_PURPLE, "Debug Mode",
         "Log keyword-to-topic matches, synonym expansions, rejection reasons",
         "CLI: --debug flag\nGUI: Debug checkbox in Matching Settings\n\nPrints per-keyword match decisions, synonym expansions, and below-threshold rejections to console."),
    ]
    card_w, card_h = Inches(3.9), Inches(4.8)
    x, y = Inches(0.7), Inches(1.35)
    for color, title, subtitle, detail in features:
        add_colored_rect(slide, x, y, card_w, card_h, C_WHITE)
        add_colored_rect(slide, x, y, card_w, Inches(0.08), color)
        add_text_box(slide, title,
                     x + Inches(0.15), y + Inches(0.2),
                     Inches(3.6), Inches(0.6),
                     font_size=14, bold=True, color=C_DARK)
        add_text_box(slide, subtitle,
                     x + Inches(0.15), y + Inches(0.9),
                     Inches(3.6), Inches(0.65),
                     font_size=12, color=C_DARK)
        add_colored_rect(slide, x + Inches(0.15), y + Inches(1.6),
                         Inches(3.55), Inches(0.03), C_LGRAY)
        add_text_box(slide, detail,
                     x + Inches(0.15), y + Inches(1.7),
                     Inches(3.6), Inches(2.8),
                     font_size=11, color=C_MUTED)
        x += card_w + Inches(0.35)


# ---------------------------------------------------------------------------
# Slide 19 — Quick Start (bookend, C_DARK)
# ---------------------------------------------------------------------------
def slide_19_quick_start(prs):
    print("Slide 19/19: Quick Start")
    slide = _blank(prs)
    add_background_fill(slide, C_DARK)
    add_text_box(slide, "Quick Start",
                 Inches(0.7), Inches(0.35), Inches(11.5), Inches(0.8),
                 font_size=32, bold=True, color=C_ACCENT)
    steps = [
        "Install:    pip install -r requirements.txt",
        "Launch:     launch_gui.bat   or   python taxonomy_matcher_gui.py",
        "Select country + browse to your two input files",
        "Click Run Matching \u2014 output opens automatically",
        "Campaign Assistant \u2192 Auto-fill \u2192 Analyse \u2192 check match rate",
        "Apply suggestions \u2192 Re-run \u2192 repeat until target reached",
    ]
    y = Inches(1.3)
    for step in steps:
        add_colored_rect(slide, Inches(0.7), y + Inches(0.05),
                         Inches(0.36), Inches(0.36), C_ACCENT)
        add_text_box(slide, "\u2713",
                     Inches(0.7), y + Inches(0.03), Inches(0.36), Inches(0.36),
                     font_size=12, bold=True, color=C_DARK, align=PP_ALIGN.CENTER)
        add_text_box(slide, step,
                     Inches(1.2), y, Inches(11.7), Inches(0.45),
                     font_size=15, color=C_WHITE)
        y += Inches(0.82)
    # Bottom strip
    add_colored_rect(slide, Inches(0), Inches(6.85), SLIDE_W, Inches(0.65), C_BLUE)
    add_text_box(slide,
                 "Output: taxonomy_match_{CC}_{timestamp}.xlsx  \u00b7  "
                 "Authoritative rate: Gap Analysis Report",
                 Inches(0.3), Inches(6.87), Inches(12.7), Inches(0.5),
                 font_size=13, color=C_WHITE, align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    prs = _new_prs()
    slide_01_title(prs)
    slide_02_problem(prs)
    slide_03_how_it_works(prs)
    slide_04_countries(prs)
    slide_05_workflow(prs)
    slide_06_input_files(prs)
    slide_07_running_matcher(prs)
    slide_08_output(prs)
    slide_09_quality(prs)
    slide_10_post_run(prs)
    slide_11_gap_analysis(prs)
    slide_12_campaign_assistant(prs)
    slide_section_improvement(prs)
    slide_14_improvement_loop(prs)
    slide_15_synonyms(prs)
    slide_16_unmapped_review(prs)
    slide_17_reports(prs)
    slide_18_advanced(prs)
    slide_19_quick_start(prs)

    out = "NL_Taxonomy_Mapper_V3_Presentation.pptx"
    prs.save(out)
    print(f"\nSaved: {out}")


if __name__ == "__main__":
    main()
